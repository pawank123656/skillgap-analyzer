"""Core logic: resume parsing/scoring, skill detection, gap analysis, roadmap, JD analysis."""
import io
import re

from data import SKILLS, ROLES

ACTION_VERBS = {
    "developed", "built", "designed", "implemented", "created", "led", "managed", "analyzed", "analysed",
    "improved", "increased", "reduced", "optimized", "optimised", "launched", "organized", "organised",
    "coordinated", "automated", "deployed", "delivered", "presented", "negotiated", "audited", "prepared",
    "reconciled", "streamlined", "achieved", "initiated", "conducted", "trained", "mentored", "generated",
    "executed", "established", "maintained", "researched", "collaborated", "supervised", "resolved",
}

SECTION_PATTERNS = {
    "education": r"\beducation\b|\bacademic",
    "skills": r"\bskills?\b|\btechnical skills\b|\bcore competenc",
    "experience": r"\bexperience\b|\binternship|\bwork history|\bemployment",
    "projects": r"\bprojects?\b|\bacademic projects\b",
    "certifications": r"\bcertifications?\b|\bcertificates?\b|\bcourses?\b",
    "summary": r"\bsummary\b|\bobjective\b|\bprofile\b|\babout me\b",
}

QUANT_RE = re.compile(
    r"\d+(?:\.\d+)?\s?%|(?:rs\.?|inr|₹|\$)\s?\d[\d,]*|"
    r"\b\d+\+?\s+(?:users|students|clients|customers|projects|members|transactions|records|leads|reports|"
    r"teams|people|participants|vendors|invoices|accounts|campaigns)\b", re.I)


# ---------------------------------------------------------------- text extraction
def extract_text(filename, data):
    name = filename.lower()
    if name.endswith(".pdf"):
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((p.extract_text() or "") for p in reader.pages)
    if name.endswith(".docx"):
        import docx
        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" ".join(c.text for c in row.cells))
        return "\n".join(parts)
    if name.endswith(".txt"):
        return data.decode("utf-8", errors="ignore")
    raise ValueError("Unsupported file type. Upload a PDF, DOCX or TXT file.")


# ---------------------------------------------------------------- skill detection
def kw_in(text, kw):
    kw = kw.lower()
    if re.fullmatch(r"[a-z0-9 ]+", kw):
        return re.search(r"\b" + re.escape(kw.strip()) + r"(?:s|es)?\b", text) is not None
    return kw in text


def detect_skills(text, course):
    """Return {skill: [matched keywords]} for the course catalog."""
    t = text.lower()
    found = {}
    for s in SKILLS[course]:
        hits = [k for k in s["kw"] if kw_in(t, k)]
        if hits:
            found[s["name"]] = hits[:4]
    return found


# ---------------------------------------------------------------- resume analysis
def analyze_resume(text, course, role):
    t = text.lower()
    words = re.findall(r"[A-Za-z][A-Za-z+#&/-]*", text)
    wc = len(words)
    reqs = ROLES[course][role]
    found = detect_skills(text, course)

    # 1. role-skill match (40)
    total_w = sum(w for _, w in reqs.values())
    got_w = sum(w for sk, (_, w) in reqs.items() if sk in found)
    skill_pts = round(40 * got_w / total_w, 1)
    missing = [sk for sk in reqs if sk not in found]

    # 2. sections (25)
    present = {k: re.search(p, t) is not None for k, p in SECTION_PATTERNS.items()}
    sec_pts = (6 * present["education"] + 5 * present["skills"] +
               8 * (present["experience"] or present["projects"]) +
               3 * present["certifications"] + 3 * present["summary"])
    missing_sections = [k for k, v in present.items() if not v]

    # 3. action verbs (10)
    verb_hits = sum(1 for w in words if w.lower() in ACTION_VERBS)
    verb_pts = round(min(verb_hits / 8, 1) * 10, 1)

    # 4. quantified achievements (15)
    quant_hits = len(QUANT_RE.findall(text))
    quant_pts = round(min(quant_hits / 3, 1) * 15, 1)

    # 5. ATS / format basics (10)
    has_email = re.search(r"[\w.+-]+@[\w-]+\.[\w.]+", text) is not None
    has_phone = re.search(r"(?:\+?\d[\d\s\-()]{8,}\d)", text) is not None
    if 300 <= wc <= 1000:
        len_pts = 4
    elif 150 <= wc < 300 or 1000 < wc <= 1400:
        len_pts = 2
    else:
        len_pts = 0
    ats_pts = 3 * has_email + 3 * has_phone + len_pts

    score = round(skill_pts + sec_pts + verb_pts + quant_pts + ats_pts)

    tips = []
    if missing:
        tips.append(f"Role keywords not found for {role}: {', '.join(missing)}. Add them only where you truly have the skill "
                    "(skills section, projects or coursework).")
    for sec in missing_sections:
        tips.append(f"Add a '{sec.title()}' section.")
    if verb_hits < 8:
        tips.append(f"Start bullet points with strong action verbs (found {verb_hits}; aim for 8+): built, led, analysed, automated...")
    if quant_hits < 3:
        tips.append(f"Add measurable results, e.g. 'reduced reporting time by 30%' or 'managed 50+ invoices' (found {quant_hits}; aim for 3+).")
    if not has_email:
        tips.append("Add a professional email address.")
    if not has_phone:
        tips.append("Add a phone number.")
    if wc < 300:
        tips.append(f"Resume looks thin ({wc} words). Add project details and responsibilities.")
    elif wc > 1000:
        tips.append(f"Resume is long ({wc} words). Freshers should keep it to one page.")
    if not tips:
        tips.append("Strong resume for this role. Tailor keywords to each job description before applying.")

    return {
        "score": score,
        "breakdown": [
            {"label": "Role skill match", "got": skill_pts, "max": 40},
            {"label": "Sections complete", "got": sec_pts, "max": 25},
            {"label": "Action verbs", "got": verb_pts, "max": 10},
            {"label": "Quantified results", "got": quant_pts, "max": 15},
            {"label": "ATS basics (contact, length)", "got": ats_pts, "max": 10},
        ],
        "skills_found": found,
        "missing_skills": missing,
        "sections": present,
        "word_count": wc,
        "verbs": verb_hits,
        "quantified": quant_hits,
        "tips": tips,
    }


# ---------------------------------------------------------------- levels + gap engine
def effective_levels(course, self_ratings, quiz_results, resume_skills):
    """Combine self-rating, quiz result and resume evidence into one level (0-5) per skill.

    - quiz taken  -> average of self rating and quiz level (quiz level = 1 + 4 * fraction correct)
    - no quiz     -> self rating as given (flagged unverified)
    - resume shows the skill -> floor of 2 (there is at least some evidence)
    """
    out = {}
    for s in SKILLS[course]:
        name = s["name"]
        selfv = float(self_ratings.get(name, 0))
        q = quiz_results.get(name)
        if q and q[1]:
            quizlv = 1 + 4 * q[0] / q[1]
            lvl = (selfv + quizlv) / 2 if selfv else quizlv * 0.6
            verified = True
        else:
            lvl, quizlv, verified = selfv, None, False
        on_resume = name in resume_skills
        if on_resume and lvl < 2:
            lvl = 2.0
        out[name] = {"self": selfv, "quiz": quizlv, "resume": on_resume, "verified": verified,
                     "level": round(lvl, 1)}
    return out


def compute_gap(reqs, levels):
    """Gap Score = sum(w * max(0, req - lvl)) / sum(w * req);  Readiness = 100 * (1 - GapScore)."""
    num = den = 0.0
    rows = []
    for skill, (req, w) in reqs.items():
        lv = levels.get(skill, {}).get("level", 0.0)
        gap = max(0.0, req - lv)
        num += w * gap
        den += w * req
        status = "Strong" if gap == 0 else ("Moderate" if gap < 2 else "Critical")
        rows.append({"skill": skill, "required": req, "weight": w, "level": lv, "gap": round(gap, 1),
                     "status": status, "priority": round(w * gap, 2)})
    gap_score = num / den if den else 0
    return {"rows": rows, "gap_score": round(gap_score, 3), "readiness": round(100 * (1 - gap_score), 1)}


def readiness_for_all_roles(course, levels):
    res = []
    for role, reqs in ROLES[course].items():
        res.append((role, compute_gap(reqs, levels)["readiness"]))
    return sorted(res, key=lambda x: -x[1])


# ---------------------------------------------------------------- roadmap
def build_roadmap(gap_rows, weeks):
    """Spread the gap skills over `weeks`, most critical first, in proportion to priority."""
    todo = sorted([r for r in gap_rows if r["gap"] > 0], key=lambda r: -r["priority"])
    if not todo:
        return []
    # at least 1 week each; drop the least important skills when weeks are fewer than skills
    todo = todo[:weeks]
    total = sum(r["priority"] for r in todo)
    alloc = {r["skill"]: max(1, round(weeks * r["priority"] / total)) for r in todo}
    while sum(alloc.values()) > weeks:                       # trim the largest block
        k = max(alloc, key=alloc.get)
        if alloc[k] <= 1:
            break
        alloc[k] -= 1
    i = 0
    while sum(alloc.values()) < weeks:                       # leftover weeks go to top priorities first
        alloc[todo[i % len(todo)]["skill"]] += 1
        i += 1
    plan, wk = [], 1
    for r in todo:
        n = alloc[r["skill"]]
        plan.append({"skill": r["skill"], "status": r["status"], "from": r["level"], "to": r["required"],
                     "start": wk, "end": wk + n - 1, "weeks": n,
                     "goal": f"Take {r['skill']} from level {r['level']} to {r['required']}: learn, practise, "
                             f"build one small project or case, then retake the quiz."})
        wk += n
    return plan


# ---------------------------------------------------------------- JD analyzer
def analyze_jd(jd_text, course, levels):
    found = detect_skills(jd_text, course)
    rows = []
    for sk in found:
        lv = levels.get(sk, {}).get("level", 0)
        rows.append({"skill": sk, "level": lv, "have": lv >= 3})
    match = round(100 * sum(r["have"] for r in rows) / len(rows)) if rows else 0
    return {"required": rows, "match": match,
            "missing": [r["skill"] for r in rows if not r["have"]]}


# ---------------------------------------------------------------- curriculum gap
def curriculum_gap(course, syllabus_text):
    covered, uncovered = [], []
    demand = {}
    for reqs in ROLES[course].values():
        for sk, (_, w) in reqs.items():
            demand[sk] = demand.get(sk, 0) + w
    t = syllabus_text.lower()
    for s in SKILLS[course]:
        hits = [k for k in s["kw"] if kw_in(t, k)]
        item = {"skill": s["name"], "demand": demand.get(s["name"], 0), "hits": hits[:4]}
        (covered if hits else uncovered).append(item)
    uncovered.sort(key=lambda x: -x["demand"])
    coverage = round(100 * len(covered) / len(SKILLS[course]))
    return {"coverage": coverage, "covered": covered, "uncovered": uncovered}
